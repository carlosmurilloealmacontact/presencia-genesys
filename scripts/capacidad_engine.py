"""
Motor de Diagnóstico y Capacidad Operativa (Requerido del Mes vs. Genesys Real).

Diseñado con base estricta en el debate gerencial de operaciones y planificación:
- Matriz Ejecutiva Panorámica: Evalúa todos los servicios de un vistazo según las 4 palancas
  (Tráfico, AHT Meta Plana, % Auxiliares vs Meta 14%, y Personas Requeridas vs Conectadas).
- Diagnóstico de Causa Raíz / Atribución de Incumplimiento.
- Vista Detallada por Franja Horaria (Lógica 30 min * Personas = Minutos Requeridos vs Disponibles).
"""

import os
import pickle
import sqlite3
from datetime import datetime, date, timedelta
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
import streamlit as st
import plotly.graph_objects as go
from config import DB_PATH
from live_engine import obtener_token_genesys

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE_FORECAST_IN = os.path.join(BASE_DIR, "../09. Intraday Forecast IN Septiembre - Latam.xlsx")
FILE_FORECAST_BO = os.path.join(BASE_DIR, "../09. Intraday Forecast BO Septiembre - Latam.xlsx")

FILE_FORECAST_DAILY = os.path.join(BASE_DIR, "../09. Daily Forecast Sept - Latam.xlsx")
CACHE_PKL_PATH = os.path.join(BASE_DIR, "../data/forecast_cache.pkl")

META_AUXILIARES_OFICIAL = 14.0  # Meta estándar de auxiliares acordada en la conversación (14%)

# Catálogo oficial maestro: nombres de servicio en mayúsculas idénticos a segments en Genesys DB
CATALOGO_SERVICIOS_SORE = {
    # Línea / Inbound Voz
    "LUA AMC": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "LUA AMC", "meta_aht": 860.0, "meta_ns": 75.0},
    "DT FFP AMC": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "DREAM TEAMS VOZ", "meta_aht": 646.0, "meta_ns": 85.0},
    "DT FFP AMC ING": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "DREAM TEAMS ENG", "meta_aht": 582.0, "meta_ns": 85.0},
    "LUA AMC ING": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "LUA AMC ING", "meta_aht": 900.0, "meta_ns": 75.0},
    "VENTAS AMC": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "Ventas AMC", "meta_aht": 780.0, "meta_ns": 75.0},
    "EQUIPAJES AMC": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "Equipajes AMC", "meta_aht": 500.0, "meta_ns": 75.0},
    "EQUIPAJES AMC ING": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "Equipajes AMC ING", "meta_aht": 491.0, "meta_ns": 75.0},
    "HVC AMC": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "HVC AMC", "meta_aht": 709.0, "meta_ns": 80.0},
    "SOPORTE LUA AMC": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "Soporte LUA AMC", "meta_aht": 311.0, "meta_ns": 75.0},
    "CORPORATE PYME": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "CORPORATE PYME", "meta_aht": 817.0, "meta_ns": 80.0},
    "AGENCIAS TARGET ES": {"tipo": "Línea / Inbound Voz", "origen": "IN", "sheet": "AGENCIAS TARGET ES", "meta_aht": 880.0, "meta_ns": 80.0},

    # Canales Digitales (Chat, WhatsApp, Redes)
    "WPP LUA AMC": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "WPP LUA AMC", "meta_aht": 1600.0, "meta_ns": 80.0},
    "WPP VENTAS AMC": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "WPP Ventas AMC", "meta_aht": 1700.0, "meta_ns": 80.0},
    "CHAT VENTAS AMC": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "CHAT Ventas AMC", "meta_aht": 1412.0, "meta_ns": 80.0},
    "WPP EQUIPAJES AMC": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "WPP EQUIPAJES AMC", "meta_aht": 1231.0, "meta_ns": 80.0},
    "RRSS AMC": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "RRSS AMC", "meta_aht": 1200.0, "meta_ns": 80.0},
    "RRSS AMC ING": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "RRSS AMC ING", "meta_aht": 1200.0, "meta_ns": 80.0},
    "RRSS PORT AMC": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "RRSS PORT AMC", "meta_aht": 1200.0, "meta_ns": 80.0},
    "AG CORPORATE CHAT": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "AGENCIAS CHAT CORPORATE", "meta_aht": 1859.0, "meta_ns": 80.0},
    "AGY N1 ESP CHAT": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "CHAT AGENCIAS ESP", "meta_aht": 1223.0, "meta_ns": 80.0, "factor_req": 0.53},
    "AGY N3 ESP CHAT": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "CHAT AGENCIAS ESP", "meta_aht": 1223.0, "meta_ns": 80.0, "factor_req": 0.47},
    "CHAT DT FFP AMC ESP": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "CHAT DREAM TEAMS ES", "meta_aht": 1103.0, "meta_ns": 80.0},
    "DREAM TEAM WP": {"tipo": "Canales Digitales", "origen": "IN", "sheet": "DREAM TEAMS WA", "meta_aht": 1272.0, "meta_ns": 80.0},

    # Back Office
    "BO LUA AMC": {"tipo": "Back Office", "origen": "BO", "sheet": "BO LUA AMC", "meta_aht": 1006.0, "meta_ns": 85.0},
    "BO EQUIPAJES AMC": {"tipo": "Back Office", "origen": "BO", "sheet": "EQUIPAJE BO", "meta_aht": 1715.0, "meta_ns": 85.0},
    "BO_CORPORATE": {"tipo": "Back Office", "origen": "BO", "sheet": "BO_CORPORATE", "meta_aht": 735.0, "meta_ns": 85.0},
    "BO AGENCIAS TARGET": {"tipo": "Back Office", "origen": "BO", "sheet": "BO AGENCIAS TARGET", "meta_aht": 735.0, "meta_ns": 85.0},
    "LATAM TRAVEL AMC": {"tipo": "Back Office", "origen": "BO", "sheet": "Latam Travel AMC", "meta_aht": 735.0, "meta_ns": 85.0},
    "BO RECLAMOS AMC": {"tipo": "Back Office", "origen": "BO", "sheet": "BO RECLAMOS AMC", "meta_aht": 2375.0, "meta_ns": 85.0},
    "BO AG LTRADE": {"tipo": "Back Office", "origen": "BO", "sheet": "BO AG LTRADE", "meta_aht": 1200.0, "meta_ns": 85.0},
    "DREAM TEAM CASOS": {"tipo": "Back Office", "origen": "BO", "sheet": "DREAM TEAM CASOS", "meta_aht": 1150.0, "meta_ns": 85.0},
}

# ── MAPEOS DE HOMOLOGACIÓN OPERATIVA WFM vs GENESYS ─────────────────────────
# 1. Homologación de presencia (de nombres en segments / nómina hacia SORE):
HOMOLOGACION_PRESENCIA_A_SORE = {
    # Agencias Voz: Niveles N1 y N3 se homologan a la exigencia de Agencias Target Voz
    "AGY N1 ESP VOZ": "AGENCIAS TARGET ES",
    "AGY N3 ESP VOZ": "AGENCIAS TARGET ES",

    # Chat Agencias: Nombre obsoleto CHAT AGENCIAS ESP redirigido a Nivel 1
    "CHAT AGENCIAS ESP": "AGY N1 ESP CHAT",

    # Back Office Reclamos: Customer Service Colombia
    "BO_CUS_COL": "BO RECLAMOS AMC",

    # Dream Team: Presencia unificada en Genesys al macro-servicio
    "DT FFP AMC": "DT FFP AMC (DREAM TEAM)",
}

# 2. Homologación de requerimiento SORE (Consolidación de streams multi-canal):
HOMOLOGACION_FORECAST_A_OPERATIVO = {
    # Dream Team: 4 streams (Voz, Chat, WhatsApp y Casos BO) gestionados por el mismo equipo multi-skill
    "CHAT DT FFP AMC ESP": "DT FFP AMC (DREAM TEAM)",
    "DREAM TEAM WP": "DT FFP AMC (DREAM TEAM)",
    "DREAM TEAM CASOS": "DT FFP AMC (DREAM TEAM)",
    "DT FFP AMC": "DT FFP AMC (DREAM TEAM)",

    # Back Office LTrade: Tarea residual absorbida por la mesa de Agencias Target
    "BO AG LTRADE": "BO AGENCIAS TARGET",
}


def obtener_servicios_raw_para_sore(servicio_sore: str) -> list[str]:
    """Retorna la lista de nombres en segments que corresponden a un servicio de SORE homologado."""
    raws = []
    for raw, mapped in HOMOLOGACION_PRESENCIA_A_SORE.items():
        if mapped == servicio_sore:
            raws.append(raw)
    if not raws:
        raws.append(servicio_sore)
    if servicio_sore not in raws and "DREAM TEAM" not in servicio_sore:
        raws.append(servicio_sore)
    return list(set(raws))



def parsear_archivos_sore_crudos() -> pd.DataFrame:
    """Parsea los libros de Excel de SORE (Inbound y Back Office) a un DataFrame unificado."""
    records = []
    archivos = {"IN": FILE_FORECAST_IN, "BO": FILE_FORECAST_BO}

    for origen, fpath in archivos.items():
        if not os.path.exists(fpath):
            continue

        wb = openpyxl.load_workbook(fpath, data_only=True)
        for srv_nombre, info in CATALOGO_SERVICIOS_SORE.items():
            if info["origen"] != origen:
                continue

            sheet_target = info["sheet"]
            if sheet_target not in wb.sheetnames:
                continue

            ws = wb[sheet_target]
            factor = float(info.get("factor_req", 1.0))
            for r in range(5, ws.max_row + 1):
                d_val = ws.cell(r, 4).value
                if not d_val:
                    continue

                f_str = d_val.strftime("%Y-%m-%d") if isinstance(d_val, datetime) else str(d_val)[:10]

                int_val = ws.cell(r, 5).value
                if isinstance(int_val, datetime) or hasattr(int_val, "strftime"):
                    int_str = int_val.strftime("%H:%M")
                else:
                    int_str = str(int_val).strip()[:5]

                traffic = ws.cell(r, 6).value or 0.0
                aht_plan = ws.cell(r, 7).value or 0.0
                asesores = ws.cell(r, 8).value or 0.0

                try:
                    traffic = float(traffic) * factor
                except Exception:
                    traffic = 0.0
                try:
                    aht_plan = float(aht_plan)
                except Exception:
                    aht_plan = 0.0
                try:
                    asesores = float(asesores) * factor
                except Exception:
                    asesores = 0.0

                records.append({
                    "servicio": srv_nombre,
                    "tipo_mundo": info["tipo"],
                    "origen": origen,
                    "fecha": f_str,
                    "intervalo": int_str,
                    "traffic_forecast": traffic,
                    "aht_forecast": aht_plan,
                    "asesores_req": asesores,
                    "minutos_req": asesores * 30.0,
                    "meta_aht_plana": info["meta_aht"],
                    "meta_ns": info["meta_ns"],
                })
        wb.close()

    df_out = pd.DataFrame(records)
    if not df_out.empty:
        # Aplicar homologación operativa de streams de requerimiento
        df_out["servicio"] = df_out["servicio"].map(lambda s: HOMOLOGACION_FORECAST_A_OPERATIVO.get(s, s))
        df_out.loc[df_out["servicio"] == "DT FFP AMC (DREAM TEAM)", "tipo_mundo"] = "Multi-canal (Voz, Chat, BO)"

        # Agrupar intervalos consolidados (por ejemplo los 4 componentes de Dream Team en cada fecha e intervalo)
        df_out = df_out.groupby(["servicio", "tipo_mundo", "fecha", "intervalo"], as_index=False).agg({
            "traffic_forecast": "sum",
            "aht_forecast": "mean",
            "asesores_req": "sum",
            "minutos_req": "sum",
            "meta_aht_plana": "mean",
            "meta_ns": "mean"
        })
    return df_out


@st.cache_data(ttl=3600, show_spinner=False)
def cargar_forecast_sore_completo(forzar_recarga: bool = False) -> pd.DataFrame:
    """
    Carga el forecast unificado desde caché ultrarrápido (pickle ~0.01s).
    Si el caché no existe o los archivos de Excel fueron modificados, lo reconstruye.
    """
    if not forzar_recarga and os.path.exists(CACHE_PKL_PATH) and os.path.getsize(CACHE_PKL_PATH) > 1000:
        cache_mtime = os.path.getmtime(CACHE_PKL_PATH)
        archivos_modificados = any(
            os.path.exists(fp) and os.path.getmtime(fp) > cache_mtime
            for fp in [FILE_FORECAST_IN, FILE_FORECAST_BO]
        )
        if not archivos_modificados:
            try:
                with open(CACHE_PKL_PATH, "rb") as f:
                    data = pickle.load(f)
                    return data
            except Exception:
                pass

    df_nuevo = parsear_archivos_sore_crudos()
    if not df_nuevo.empty:
        try:
            os.makedirs(os.path.dirname(CACHE_PKL_PATH), exist_ok=True)
            tmp_cache = CACHE_PKL_PATH + ".tmp"
            with open(tmp_cache, "wb") as f:
                pickle.dump(df_nuevo, f, protocol=pickle.HIGHEST_PROTOCOL)
                f.flush()
                os.fsync(f.fileno())
            os.replace(tmp_cache, CACHE_PKL_PATH)
        except Exception:
            pass
    return df_nuevo


def cargar_presencia_resumen_rango(fecha_desde: str, fecha_hasta: str, num_dias: int = 1) -> pd.DataFrame:
    """
    Agrupa los minutos de presencia real de Genesys para un rango de fechas (o un solo día).
    Reconoce de forma inteligente los estados productivos según el tipo de servicio:
    - Inbound / Voz: Available y On Queue son productivos.
    - Back Office y Dream Team: Casos Backoffice, Available y On Queue son productivos.
    Aplica homologación operativa de servicios para alinear con SORE.
    """
    real_db_path = Path(__file__).parent / DB_PATH
    if not os.path.exists(real_db_path):
        return pd.DataFrame()

    conn = sqlite3.connect(real_db_path)
    query = """
        SELECT 
            UPPER(TRIM(servicio)) as servicio,
            COUNT(DISTINCT agente_id) as asesores_conectados,
            SUM(duracion_min) as min_total,
            SUM(CASE WHEN presence_label NOT IN ('Offline') THEN duracion_min ELSE 0 END) as min_conectado,
            SUM(CASE 
                WHEN UPPER(TRIM(servicio)) LIKE '%BO%' 
                  OR UPPER(TRIM(servicio)) LIKE '%BACKOFFICE%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DREAM%' 
                  OR UPPER(TRIM(servicio)) LIKE '%AGY%CHAT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%CHAT%AGENCIA%'
                  OR UPPER(TRIM(servicio)) LIKE '%AGENCIA%CHAT%'
                  OR UPPER(TRIM(servicio)) LIKE '%AG CORPORATE CHAT%' THEN
                    CASE WHEN presence_label IN ('Available', 'On Queue', 'Conectado', 'Casos Backoffice', 'Gestión sin Contacto', 'Gestion sin Contacto') 
                              OR UPPER(presence_label) LIKE '%BACKOFFICE%' 
                              OR UPPER(presence_label) LIKE '%GESTION%' 
                              OR UPPER(presence_label) LIKE '%GESTIÓN%' 
                         THEN duracion_min ELSE 0 END
                ELSE
                    CASE WHEN presence_label IN ('Available', 'On Queue', 'Conectado') THEN duracion_min ELSE 0 END
            END) as min_disponible,
            SUM(CASE 
                WHEN UPPER(TRIM(servicio)) LIKE '%BO%' 
                  OR UPPER(TRIM(servicio)) LIKE '%BACKOFFICE%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DREAM%' 
                  OR UPPER(TRIM(servicio)) LIKE '%AGY%CHAT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%CHAT%AGENCIA%'
                  OR UPPER(TRIM(servicio)) LIKE '%AGENCIA%CHAT%'
                  OR UPPER(TRIM(servicio)) LIKE '%AG CORPORATE CHAT%' THEN
                    CASE WHEN presence_label NOT IN ('Offline', 'Available', 'On Queue', 'Conectado', 'Casos Backoffice', 'Gestión sin Contacto', 'Gestion sin Contacto') 
                              AND UPPER(presence_label) NOT LIKE '%BACKOFFICE%' 
                              AND UPPER(presence_label) NOT LIKE '%GESTION%' 
                              AND UPPER(presence_label) NOT LIKE '%GESTIÓN%' 
                         THEN duracion_min ELSE 0 END
                ELSE
                    CASE WHEN presence_label NOT IN ('Offline', 'Available', 'On Queue', 'Conectado') THEN duracion_min ELSE 0 END
            END) as min_pausas
        FROM segments
        WHERE fecha >= ? AND fecha <= ? AND servicio IS NOT NULL AND servicio != ''
        GROUP BY UPPER(TRIM(servicio))
    """
    df_pres = pd.read_sql(query, conn, params=(fecha_desde, fecha_hasta))
    conn.close()

    if not df_pres.empty:
        # Homologar nombres de servicio de nómina a la nomenclatura oficial de dimensionamiento
        df_pres["servicio"] = df_pres["servicio"].map(lambda s: HOMOLOGACION_PRESENCIA_A_SORE.get(s, s))
        df_pres = df_pres.groupby("servicio", as_index=False).agg({
            "asesores_conectados": "sum",
            "min_total": "sum",
            "min_conectado": "sum",
            "min_disponible": "sum",
            "min_pausas": "sum"
        })

        df_pres["pct_auxiliares_real"] = df_pres.apply(
            lambda r: (r["min_pausas"] / r["min_conectado"] * 100.0) if r["min_conectado"] > 0 else 0.0, axis=1
        ).round(1)
        # FTEs equivalentes promedio diario en base a jornada estándar de 8 horas (480 minutos * num_dias)
        dias_div = max(1, num_dias)
        df_pres["fte_reales_conectados"] = (df_pres["min_conectado"] / (480.0 * dias_div)).round(1)
        df_pres["fte_reales_disponibles"] = (df_pres["min_disponible"] / (480.0 * dias_div)).round(1)

    return df_pres


def cargar_presencia_resumen_dia(fecha_str: str) -> pd.DataFrame:
    """Compatibilidad: Agrupa minutos de presencia real para un solo día."""
    return cargar_presencia_resumen_rango(fecha_str, fecha_str, num_dias=1)


def calcular_evolucion_diaria_servicio(
    fecha_desde: str, fecha_hasta: str, servicio_sel: str, df_fore_all: pd.DataFrame
) -> pd.DataFrame:
    """
    Calcula la serie temporal día a día de un servicio:
    FTE Requerido vs FTE Conectado vs FTE Disponible y % de Capacidad.
    """
    real_db_path = Path(__file__).parent / DB_PATH
    if not os.path.exists(real_db_path):
        return pd.DataFrame()

    raws = obtener_servicios_raw_para_sore(servicio_sel)
    placeholders = ", ".join("?" for _ in raws)

    conn = sqlite3.connect(real_db_path)
    query = f"""
        SELECT 
            fecha,
            COUNT(DISTINCT agente_id) as asesores_conectados,
            SUM(duracion_min) as min_total,
            SUM(CASE WHEN presence_label NOT IN ('Offline') THEN duracion_min ELSE 0 END) as min_conectado,
            SUM(CASE 
                WHEN UPPER(TRIM(servicio)) LIKE '%BO%' 
                  OR UPPER(TRIM(servicio)) LIKE '%BACKOFFICE%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DREAM%' 
                  OR UPPER(TRIM(servicio)) LIKE '%AGY%CHAT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%CHAT%AGENCIA%'
                  OR UPPER(TRIM(servicio)) LIKE '%AGENCIA%CHAT%'
                  OR UPPER(TRIM(servicio)) LIKE '%AG CORPORATE CHAT%' THEN
                    CASE WHEN presence_label IN ('Available', 'On Queue', 'Conectado', 'Casos Backoffice', 'Gestión sin Contacto', 'Gestion sin Contacto') 
                              OR UPPER(presence_label) LIKE '%BACKOFFICE%' 
                              OR UPPER(presence_label) LIKE '%GESTION%' 
                              OR UPPER(presence_label) LIKE '%GESTIÓN%' 
                         THEN duracion_min ELSE 0 END
                ELSE
                    CASE WHEN presence_label IN ('Available', 'On Queue', 'Conectado') THEN duracion_min ELSE 0 END
            END) as min_disponible,
            SUM(CASE 
                WHEN UPPER(TRIM(servicio)) LIKE '%BO%' 
                  OR UPPER(TRIM(servicio)) LIKE '%BACKOFFICE%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%DREAM%' 
                  OR UPPER(TRIM(servicio)) LIKE '%AGY%CHAT%' 
                  OR UPPER(TRIM(servicio)) LIKE '%CHAT%AGENCIA%'
                  OR UPPER(TRIM(servicio)) LIKE '%AGENCIA%CHAT%'
                  OR UPPER(TRIM(servicio)) LIKE '%AG CORPORATE CHAT%' THEN
                    CASE WHEN presence_label NOT IN ('Offline', 'Available', 'On Queue', 'Conectado', 'Casos Backoffice', 'Gestión sin Contacto', 'Gestion sin Contacto') 
                              AND UPPER(presence_label) NOT LIKE '%BACKOFFICE%' 
                              AND UPPER(presence_label) NOT LIKE '%GESTION%' 
                              AND UPPER(presence_label) NOT LIKE '%GESTIÓN%' 
                         THEN duracion_min ELSE 0 END
                ELSE
                    CASE WHEN presence_label NOT IN ('Offline', 'Available', 'On Queue', 'Conectado') THEN duracion_min ELSE 0 END
            END) as min_pausas
        FROM segments
        WHERE fecha >= ? AND fecha <= ? AND UPPER(TRIM(servicio)) IN ({placeholders})
        GROUP BY fecha
    """
    params = (fecha_desde, fecha_hasta, *[r.upper().strip() for r in raws])
    df_pres = pd.read_sql(query, conn, params=params)
    conn.close()

    sub_f = df_fore_all[
        (df_fore_all["fecha"] >= fecha_desde)
        & (df_fore_all["fecha"] <= fecha_hasta)
        & (df_fore_all["servicio"] == servicio_sel)
    ]
    df_fore_dia = sub_f.groupby("fecha").agg({
        "traffic_forecast": "sum",
        "minutos_req": "sum"
    }).reset_index()

    if df_fore_dia.empty and df_pres.empty:
        return pd.DataFrame()

    merged = pd.merge(df_fore_dia, df_pres, on="fecha", how="outer").fillna(0.0).sort_values(by="fecha")
    merged["fte_req"] = (merged["minutos_req"] / 480.0).round(1)
    merged["fte_con"] = (merged["min_conectado"] / 480.0).round(1)
    merged["fte_disp"] = (merged["min_disponible"] / 480.0).round(1)
    merged["brecha_fte"] = (merged["fte_con"] - merged["fte_req"]).round(1)
    merged["pct_aux"] = merged.apply(
        lambda r: (r["min_pausas"] / r["min_conectado"] * 100.0) if r["min_conectado"] > 0 else 0.0, axis=1
    ).round(1)
    merged["pct_capacidad"] = merged.apply(
        lambda r: (r["min_disponible"] / r["minutos_req"] * 100.0) if r["minutos_req"] > 0 else 100.0, axis=1
    ).round(1)

    return merged


def calcular_capacidad_intervalos_real(fecha_str: str, servicio_sel: str) -> pd.DataFrame:
    """
    Calcula la presencia real de Genesys para cada uno de los 48 intervalos de 30 min.
    Aplica la distinción de estados productivos de Back Office vs Inbound.
    """
    real_db_path = Path(__file__).parent / DB_PATH
    if not os.path.exists(real_db_path):
        return pd.DataFrame()

    raws = obtener_servicios_raw_para_sore(servicio_sel)
    placeholders = ", ".join("?" for _ in raws)

    conn = sqlite3.connect(real_db_path)
    query = f"""
        SELECT agente, presence_label, system_presence, inicio, fin, duracion_min
        FROM segments
        WHERE fecha = ? AND UPPER(TRIM(servicio)) IN ({placeholders})
    """
    params = (fecha_str, *[r.upper().strip() for r in raws])
    df_seg = pd.read_sql(query, conn, params=params)
    conn.close()

    if df_seg.empty:
        return pd.DataFrame()

    df_seg["dt_ini"] = pd.to_datetime(df_seg["inicio"])
    df_seg["dt_fin"] = pd.to_datetime(df_seg["fin"]).fillna(
        df_seg["dt_ini"] + pd.to_timedelta(df_seg["duracion_min"], unit="m")
    )

    es_bo_o_chat_agy = any(k in servicio_sel.upper() for k in ["BO", "BACKOFFICE", "DREAM", "DT", "AGY", "AGENCIA"])
    estados_productivos = ["Available", "On Queue", "Conectado", "Casos Backoffice", "Gestión sin Contacto", "Gestion sin Contacto"] if es_bo_o_chat_agy else ["Available", "On Queue", "Conectado"]

    day_dt = pd.to_datetime(fecha_str)
    intervalos = []
    for i in range(48):
        t_start = day_dt + pd.Timedelta(minutes=i * 30)
        t_end = t_start + pd.Timedelta(minutes=30)
        int_label = t_start.strftime("%H:%M")

        ov_s = df_seg["dt_ini"].clip(lower=t_start)
        ov_e = df_seg["dt_fin"].clip(upper=t_end)
        ov_min = (ov_e - ov_s).dt.total_seconds() / 60.0
        ov_min = ov_min.clip(lower=0)

        sub = df_seg[ov_min > 0].copy()
        sub["w_min"] = ov_min[ov_min > 0]

        if not sub.empty:
            min_con = sub[sub["presence_label"] != "Offline"]["w_min"].sum()
            if es_bo_o_chat_agy:
                is_prod = sub["presence_label"].isin(estados_productivos) | sub["presence_label"].astype(str).str.upper().str.contains("BACKOFFICE|GESTION|GESTIÓN", na=False)
            else:
                is_prod = sub["presence_label"].isin(estados_productivos)
            min_disp = sub[is_prod]["w_min"].sum()
            min_pau = sub[(sub["presence_label"] != "Offline") & (~is_prod)]["w_min"].sum()
        else:
            min_con = 0.0
            min_disp = 0.0
            min_pau = 0.0

        intervalos.append({
            "intervalo": int_label,
            "min_conectado": min_con,
            "min_disponible": min_disp,
            "min_pausas": min_pau,
            "fte_conectado": min_con / 30.0,
            "fte_disponible": min_disp / 30.0,
        })

    return pd.DataFrame(intervalos)


def diagnosticar_causa_raiz(gap_personas: float, aux_real: float, aux_meta: float, pct_capacidad: float) -> str:
    """
    Atribuye con objetividad operativa la causa raíz por la cual un servicio no cumplió su capacidad.
    """
    if pct_capacidad >= 95.0:
        if gap_personas >= 0 and aux_real <= aux_meta:
            return "🟢 Capacidad Óptima"
        return "🟢 Capacidad Cumplida"

    causas = []
    if gap_personas < -1.5:
        causas.append(f"Déficit Personas ({gap_personas:+.1f} FTE)")
    if aux_real > (aux_meta + 2.0):
        causas.append(f"Fuga Auxiliares ({aux_real:.1f}% > {aux_meta:.0f}%)")

    if not causas:
        if pct_capacidad >= 85.0:
            return "🟡 Desviación Leve (Aceptable)"
        return "🟡 Dilución Operativa Intradía"

    return "🔴 " + " • ".join(causas)


MAPA_GTR_A_SORE = {
    "LUA AMC": "LUA AMC",
    "LUA AMC ING": "LUA AMC ING",
    "Soporte LUA AMC": "SOPORTE LUA AMC",
    "HVC AMC": "HVC AMC",
    "Ventas AMC": "VENTAS AMC",
    "Equipajes AMC": "EQUIPAJES AMC",
    "Equipajes AMC ING": "EQUIPAJES AMC ING",
    "WPP LUA AMC": "WPP LUA AMC",
    "WPP VENTAS AMC": "WPP VENTAS AMC",
    "CHAT VENTAS AMC": "CHAT VENTAS AMC",
    "WPP EQUIPAJES AMC": "WPP EQUIPAJES AMC",
    "RRSS AMC": "RRSS AMC",
    "DT FFP AMC ING": "DT FFP AMC ING",
    "DT FFP AMC": "DT FFP AMC (DREAM TEAM)",
    "CHAT DT FFP AMC ESP": "DT FFP AMC (DREAM TEAM)",
    "DREAM TEAM WP": "DT FFP AMC (DREAM TEAM)",
    "GSS NDC Agencias": "AGENCIAS TARGET ES",
    "GSS Operacional Agencias": "AGENCIAS TARGET ES",
    "TRAVEL WP AMC": "TRAVEL WP AMC",
}


@st.cache_data(ttl=120, show_spinner=False)
def obtener_metricas_gtr_rango(fecha_desde: str, fecha_hasta: str) -> pd.DataFrame:
    """
    Consulta las métricas agregadas de colas en Genesys Cloud para un rango de fechas (o un día)
    y las homologa a nivel de servicio SORE (Tráfico real, AHT real, NS 80/20 real).
    """
    token = obtener_token_genesys()
    if not token:
        return pd.DataFrame()
    try:
        from gtr_engine import obtener_metricas_gtr_api
        df_metrics, err, _ = obtener_metricas_gtr_api(token, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)
        if df_metrics.empty:
            return pd.DataFrame()

        def _mapear(row):
            srv = row.get("servicio", "")
            if srv in MAPA_GTR_A_SORE:
                return MAPA_GTR_A_SORE[srv]
            srv_u = str(srv).upper()
            canal = row.get("canal", "VOZ")
            if "AGENCIA" in srv_u or "AGY" in srv_u:
                if canal == "CHAT":
                    if "N3" in srv_u or "NIVEL 3" in srv_u or "NIVEL3" in srv_u:
                        return "AGY N3 ESP CHAT"
                    return "AGY N1 ESP CHAT"
                return "AGENCIAS TARGET ES"
            return srv_u

        df_metrics["servicio_sore"] = df_metrics.apply(_mapear, axis=1)
        agg = df_metrics.groupby("servicio_sore").agg({
            "nOffered": "sum",
            "tAnswered_count": "sum",
            "tHandle_sum": "sum",
            "tHandle_count": "sum",
            "sl_numerator": "sum",
            "sl_denominator": "sum"
        }).reset_index()

        agg["trafico_real"] = agg["nOffered"]
        agg["aht_real_seg"] = agg.apply(
            lambda r: int(round((r["tHandle_sum"] / r["tHandle_count"]) / 1000.0)) if r["tHandle_count"] > 0 else np.nan,
            axis=1
        )
        agg["ns_real"] = agg.apply(
            lambda r: round((r["sl_numerator"] / r["sl_denominator"] * 100.0), 1) if r["sl_denominator"] > 0 else np.nan,
            axis=1
        )
        return agg.rename(columns={"servicio_sore": "servicio"})[[
            "servicio", "trafico_real", "aht_real_seg", "ns_real", "tHandle_sum", "tHandle_count", "sl_numerator", "sl_denominator"
        ]]
    except Exception:
        return pd.DataFrame()


@st.cache_data(ttl=120, show_spinner=False)
def obtener_metricas_servicio_intradia(fecha_sel: str, srv_detalle: str) -> pd.DataFrame:
    """
    Obtiene las métricas de GTR de Genesys Cloud (Tráfico real recibido, AHT real, NS real %)
    por cada intervalo de 30 min para un servicio específico.
    """
    token = obtener_token_genesys()
    if not token:
        return pd.DataFrame()
    try:
        from gtr_engine import obtener_metricas_gtr_api
        df_metrics, err, _ = obtener_metricas_gtr_api(token, fecha_desde=fecha_sel, fecha_hasta=fecha_sel)
        if df_metrics.empty:
            return pd.DataFrame()

        srv_u = srv_detalle.upper()
        if "DREAM TEAM" in srv_u or "DT FFP" in srv_u:
            sub = df_metrics[df_metrics["servicio"].isin(["DT FFP AMC", "CHAT DT FFP AMC ESP", "DREAM TEAM WP"])]
        elif ("AGENCIAS" in srv_u or "AGY" in srv_u) and "CHAT" not in srv_u:
            sub = df_metrics[df_metrics["servicio"].str.contains("Agencias|AGY", case=False, na=False) & (df_metrics["canal"] == "VOZ")]
        elif ("AGENCIAS" in srv_u or "AGY" in srv_u) and "CHAT" in srv_u:
            if "N3" in srv_u or "NIVEL 3" in srv_u or "NIVEL3" in srv_u:
                sub = df_metrics[df_metrics["servicio"].str.contains("N3|Nivel 3|Nivel3", case=False, na=False) & (df_metrics["canal"] == "CHAT")]
                if sub.empty:
                    sub = df_metrics[df_metrics["servicio"].str.contains("Agencias|AGY", case=False, na=False) & (df_metrics["canal"] == "CHAT")]
            else:
                sub = df_metrics[df_metrics["servicio"].str.contains("N1|Nivel 1|Nivel1|Expert", case=False, na=False) & (df_metrics["canal"] == "CHAT")]
                if sub.empty:
                    sub = df_metrics[df_metrics["servicio"].str.contains("Agencias|AGY", case=False, na=False) & (df_metrics["canal"] == "CHAT")]
        else:
            sub = df_metrics[df_metrics["servicio"].str.upper() == srv_u]

        if sub.empty:
            return pd.DataFrame()

        agg = sub.groupby("intervalo").agg({
            "nOffered": "sum",
            "tAnswered_count": "sum",
            "tHandle_sum": "sum",
            "tHandle_count": "sum",
            "sl_numerator": "sum",
            "sl_denominator": "sum"
        }).reset_index()

        agg["trafico_real"] = agg["nOffered"]
        agg["aht_real_seg"] = agg.apply(
            lambda r: int(round((r["tHandle_sum"] / r["tHandle_count"]) / 1000.0)) if r["tHandle_count"] > 0 else np.nan,
            axis=1
        )
        agg["ns_real"] = agg.apply(
            lambda r: round((r["sl_numerator"] / r["sl_denominator"] * 100.0), 1) if r["sl_denominator"] > 0 else np.nan,
            axis=1
        )
        return agg[["intervalo", "trafico_real", "aht_real_seg", "ns_real"]]
    except Exception:
        return pd.DataFrame()


def _render_curva_y_tabla_intradia(fecha_sel: str, srv_detalle: str, df_fore_all: pd.DataFrame):
    """
    Renderiza la gráfica de 48 intervalos de 30 min y la tabla detallada
    con FTEs, Demanda/Tráfico real vs plan, AHT real vs meta, NS 80/20 y Capacidad Neta.
    """
    df_fore_dia = df_fore_all[df_fore_all["fecha"] == fecha_sel].copy()
    sub_f_int = df_fore_dia[df_fore_dia["servicio"] == srv_detalle].copy()
    sub_r_int = calcular_capacidad_intervalos_real(fecha_sel, srv_detalle)

    if sub_r_int.empty:
        sub_r_int = pd.DataFrame([
            {
                "intervalo": int_lbl,
                "min_conectado": 0.0,
                "min_disponible": 0.0,
                "min_pausas": 0.0,
                "fte_conectado": 0.0,
                "fte_disponible": 0.0,
            }
            for int_lbl in sub_f_int["intervalo"].tolist()
        ])

    merged_int = pd.merge(sub_f_int, sub_r_int, on="intervalo", how="left").fillna(0.0)

    # Cruzar con métricas GTR en vivo / históricas
    df_gtr_int = obtener_metricas_servicio_intradia(fecha_sel, srv_detalle)
    if not df_gtr_int.empty:
        merged_int = pd.merge(merged_int, df_gtr_int, on="intervalo", how="left")
    else:
        merged_int["trafico_real"] = np.nan
        merged_int["aht_real_seg"] = np.nan
        merged_int["ns_real"] = np.nan

    fig_int = go.Figure()
    fig_int.add_trace(go.Scatter(
        x=merged_int["intervalo"],
        y=merged_int["asesores_req"],
        mode="lines+markers",
        name="1. Requerido del Mes (FTEs)",
        line=dict(color="#f59e0b", width=3, dash="dash")
    ))
    fig_int.add_trace(go.Scatter(
        x=merged_int["intervalo"],
        y=merged_int["fte_conectado"],
        mode="lines",
        name="2. Conectados Totales (FTEs)",
        line=dict(color="#94a3b8", width=2)
    ))
    fig_int.add_trace(go.Bar(
        x=merged_int["intervalo"],
        y=merged_int["fte_disponible"],
        name="3. Disponible Efectivo (FTEs)",
        marker_color="#2563eb",
        opacity=0.75
    ))

    fig_int.update_layout(
        title=f"Curva Intradía de Cobertura y Capacidad — {srv_detalle} ({fecha_sel})",
        xaxis=dict(title="Intervalo de 30 min", tickangle=-45),
        yaxis=dict(title="Equivalente de Asesores (FTEs)"),
        hovermode="x unified",
        margin=dict(l=20, r=20, t=40, b=20),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
    )
    st.plotly_chart(fig_int, use_container_width=True)

    with st.expander(f"📋 Ver Tabla Detallada Intervalo a Intervalo ({srv_detalle} - {fecha_sel})", expanded=False):
        c_tog, _ = st.columns([3, 2])
        with c_tog:
            mostrar_minutos = st.checkbox(
                "🔍 Ver columnas de minutos brutos (Min. Requeridos, Min. Pausas, Min. Disponibles)",
                value=False,
                key=f"cap_toggle_min_{srv_detalle}_{fecha_sel}"
            )

        df_calc = merged_int.copy()
        df_calc["Brecha FTE"] = (df_calc["fte_conectado"] - df_calc["asesores_req"]).round(2)
        df_calc["% Desv Tráfico"] = df_calc.apply(
            lambda r: round(((r["trafico_real"] - r["traffic_forecast"]) / r["traffic_forecast"] * 100.0), 1)
            if pd.notna(r.get("trafico_real")) and r["traffic_forecast"] > 0 else np.nan,
            axis=1
        )
        meta_plana_val = int(round(sub_f_int["meta_aht_plana"].iloc[0])) if not sub_f_int.empty and "meta_aht_plana" in sub_f_int.columns else 800
        df_calc["AHT Plan"] = meta_plana_val
        df_calc["% Capacidad"] = df_calc.apply(
            lambda r: (r["min_disponible"] / r["minutos_req"] * 100.0) if r["minutos_req"] > 0 else 100.0,
            axis=1
        ).round(1)

        cols_base = [
            "intervalo", "asesores_req", "fte_conectado", "fte_disponible", "Brecha FTE",
            "traffic_forecast", "trafico_real", "% Desv Tráfico",
            "AHT Plan", "aht_real_seg", "ns_real", "% Capacidad"
        ]
        if mostrar_minutos:
            cols_base.extend(["minutos_req", "min_conectado", "min_pausas", "min_disponible"])

        df_mostrar_int = df_calc[cols_base].rename(columns={
            "intervalo": "Intervalo",
            "asesores_req": "FTE Req",
            "fte_conectado": "FTE Con",
            "fte_disponible": "FTE Disp",
            "traffic_forecast": "Tráfico Plan",
            "trafico_real": "Tráfico Real",
            "aht_real_seg": "AHT Real (s)",
            "ns_real": "% NS",
            "minutos_req": "Min. Requeridos",
            "min_conectado": "Min. Conectados",
            "min_pausas": "Min. Pausas",
            "min_disponible": "Min. Disponibles",
        })

        # Estilos condicionales ejecutivos
        def c_ns(val):
            if pd.isna(val): return ""
            if val >= 80.0: return "background-color: rgba(16, 185, 129, 0.20); color: #10b981; font-weight: 700;"
            elif val >= 70.0: return "background-color: rgba(245, 158, 11, 0.20); color: #f59e0b; font-weight: 700;"
            return "background-color: rgba(239, 68, 68, 0.20); color: #ef4444; font-weight: 700;"

        def c_desv(val):
            if pd.isna(val): return ""
            if val <= 5.0: return "background-color: rgba(16, 185, 129, 0.15); color: #10b981; font-weight: 600;"
            elif val <= 15.0: return "background-color: rgba(245, 158, 11, 0.15); color: #f59e0b; font-weight: 600;"
            return "background-color: rgba(239, 68, 68, 0.20); color: #ef4444; font-weight: 700;"

        def c_gap(val):
            if pd.isna(val): return ""
            if val >= 0: return "background-color: rgba(16, 185, 129, 0.15); color: #10b981; font-weight: 700;"
            elif val >= -1.0: return "background-color: rgba(245, 158, 11, 0.15); color: #f59e0b; font-weight: 700;"
            return "background-color: rgba(239, 68, 68, 0.20); color: #ef4444; font-weight: 700;"

        def c_cap(val):
            if pd.isna(val): return ""
            if val >= 95.0: return "background-color: rgba(16, 185, 129, 0.15); color: #10b981; font-weight: 700;"
            elif val >= 85.0: return "background-color: rgba(245, 158, 11, 0.15); color: #f59e0b; font-weight: 700;"
            return "background-color: rgba(239, 68, 68, 0.20); color: #ef4444; font-weight: 700;"

        def c_aht(val):
            if pd.isna(val) or val == 0: return ""
            if val <= meta_plana_val: return "background-color: rgba(16, 185, 129, 0.15); color: #10b981; font-weight: 600;"
            elif val <= meta_plana_val * 1.10: return "background-color: rgba(245, 158, 11, 0.15); color: #f59e0b; font-weight: 600;"
            return "background-color: rgba(239, 68, 68, 0.20); color: #ef4444; font-weight: 700;"

        styler_int = (
            df_mostrar_int.style
            .map(c_ns, subset=["% NS"])
            .map(c_desv, subset=["% Desv Tráfico"])
            .map(c_gap, subset=["Brecha FTE"])
            .map(c_cap, subset=["% Capacidad"])
            .map(c_aht, subset=["AHT Real (s)"])
        )

        col_configs = {
            "Intervalo": st.column_config.TextColumn(width="small"),
            "FTE Req": st.column_config.NumberColumn("FTE Req", format="%.2f"),
            "FTE Con": st.column_config.NumberColumn("FTE Con", format="%.2f"),
            "FTE Disp": st.column_config.NumberColumn("FTE Disp", format="%.2f"),
            "Brecha FTE": st.column_config.NumberColumn("Brecha FTE", format="%+.2f"),
            "Tráfico Plan": st.column_config.NumberColumn("Tráfico Plan", format="%.1f"),
            "Tráfico Real": st.column_config.NumberColumn("Tráfico Real", format="%.0f"),
            "% Desv Tráfico": st.column_config.NumberColumn("% Desv Tráfico", format="%+.1f%%"),
            "AHT Plan": st.column_config.NumberColumn("AHT Plan", format="%d s"),
            "AHT Real (s)": st.column_config.NumberColumn("AHT Real", format="%.0f s"),
            "% NS": st.column_config.NumberColumn("% NS", format="%.1f%%"),
            "% Capacidad": st.column_config.NumberColumn("% Capacidad", format="%.1f%%"),
        }
        if mostrar_minutos:
            col_configs.update({
                "Min. Requeridos": st.column_config.NumberColumn(format="%.1f m"),
                "Min. Conectados": st.column_config.NumberColumn(format="%.1f m"),
                "Min. Pausas": st.column_config.NumberColumn(format="%.1f m"),
                "Min. Disponibles": st.column_config.NumberColumn(format="%.1f m"),
            })

        st.dataframe(
            styler_int,
            use_container_width=True,
            hide_index=True,
            column_config=col_configs
        )


def render_tab_capacidad(agentes_map: dict):
    """
    Renderiza la Matriz Ejecutiva Panorámica de Capacidad y el Desglose Intradía.
    """
    st.markdown("### 🧭 Matriz Ejecutiva de Capacidad y Diagnóstico Operativo")
    st.caption(
        "Herramienta gerencial de contraste: Compara la base del requerido del mes "
        "frente a la ejecución real de presencia en Genesys, identificando la causa raíz de las brechas de servicio."
    )

    df_fore_all = cargar_forecast_sore_completo()
    if df_fore_all.empty:
        st.error("⚠️ No se encontraron los archivos base del requerido del mes en la raíz del proyecto.")
        return

    # Selector de fecha disponible (default al día más reciente con registros)
    fechas_disp = sorted(df_fore_all["fecha"].unique().tolist())
    real_db_path = Path(__file__).parent / DB_PATH
    fecha_default_idx = len(fechas_disp) - 1
    if os.path.exists(real_db_path):
        try:
            conn = sqlite3.connect(real_db_path)
            c = conn.cursor()
            c.execute("SELECT MAX(fecha) FROM segments WHERE fecha IS NOT NULL AND fecha >= '2026-09-01'")
            max_f = c.fetchone()[0]
            conn.close()
            if max_f and max_f in fechas_disp:
                fecha_default_idx = fechas_disp.index(max_f)
        except Exception:
            pass

    # Barra superior de controles
    col_t1, col_t2, col_f2, col_f3, col_btn = st.columns([1.3, 2.2, 1.8, 3.2, 1.0])
    with col_t1:
        tipo_temporalidad = st.radio(
            "Temporalidad:",
            ["📅 Día", "📊 Rango"],
            horizontal=True,
            key="cap_tipo_temporalidad"
        )
    with col_t2:
        if tipo_temporalidad == "📅 Día":
            fecha_sel = st.selectbox("Fecha de Evaluación:", fechas_disp, index=fecha_default_idx, key="cap_fecha_sel")
            fecha_desde = fecha_sel
            fecha_hasta = fecha_sel
            num_dias = 1
            texto_periodo = f"Día: <b>{fecha_sel}</b>"
        else:
            min_d = datetime.strptime(fechas_disp[0], "%Y-%m-%d").date()
            max_d = datetime.strptime(fechas_disp[-1], "%Y-%m-%d").date()
            default_start = min_d
            default_end = datetime.strptime(fechas_disp[fecha_default_idx], "%Y-%m-%d").date()
            rango_pick = st.date_input(
                "Rango de Fechas:",
                value=(default_start, default_end),
                min_value=min_d,
                max_value=max_d,
                key="cap_rango_fechas"
            )
            if isinstance(rango_pick, (tuple, list)) and len(rango_pick) == 2:
                fecha_desde = rango_pick[0].strftime("%Y-%m-%d")
                fecha_hasta = rango_pick[1].strftime("%Y-%m-%d")
            elif isinstance(rango_pick, (tuple, list)) and len(rango_pick) == 1:
                fecha_desde = rango_pick[0].strftime("%Y-%m-%d")
                fecha_hasta = fecha_desde
            else:
                fecha_desde = default_start.strftime("%Y-%m-%d")
                fecha_hasta = default_end.strftime("%Y-%m-%d")

            if fecha_desde > fecha_hasta:
                fecha_desde, fecha_hasta = fecha_hasta, fecha_desde

            dias_en_rango = [f for f in fechas_disp if fecha_desde <= f <= fecha_hasta]
            num_dias = max(1, len(dias_en_rango))
            texto_periodo = f"Rango: <b>{fecha_desde}</b> a <b>{fecha_hasta}</b> ({num_dias}d)"

    with col_f2:
        filtro_mundo = st.selectbox(
            "🌐 Operación / Canal:",
            ["Todos los Servicios", "📞 Línea / Inbound Voz", "💬 Canales Digitales", "📂 Back Office"],
            index=0
        )
    with col_f3:
        st.markdown(
            f"""
            <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 7px 12px; margin-top: 5px;">
                <span style="font-size: 11px; color: #64748b;">Parámetros del Comité · {texto_periodo}:</span><br>
                <span style="font-size: 12px; font-weight: 600; color: #0f172a;">
                    Meta Auxiliares: <b>{META_AUXILIARES_OFICIAL:.0f}%</b> (86% Disp.) · Meta AHT: <b>Mes</b> · Base FTE: <b>8h (480m)</b>
                </span>
            </div>
            """,
            unsafe_allow_html=True
        )
    with col_btn:
        st.markdown("<div style='margin-top: 24px;'></div>", unsafe_allow_html=True)
        if st.button("🔄 Recargar", help="Fuerza la lectura fresca de los archivos base del requerido del mes"):
            cargar_forecast_sore_completo(forzar_recarga=True)
            st.rerun()

    # 1. Consolidar dimensionamiento del periodo
    df_fore_rango = df_fore_all[(df_fore_all["fecha"] >= fecha_desde) & (df_fore_all["fecha"] <= fecha_hasta)].copy()
    if df_fore_rango.empty:
        st.warning(f"No hay registros de dimensionamiento para las fechas seleccionadas ({fecha_desde} a {fecha_hasta}).")
        return

    # 2. Cargar presencia real de Genesys
    df_pres_rango = cargar_presencia_resumen_rango(fecha_desde, fecha_hasta, num_dias=num_dias)

    # 3. Cargar métricas de Genesys GTR (Tráfico real, AHT real, NS 80/20 real)
    df_gtr_rango = obtener_metricas_gtr_rango(fecha_desde, fecha_hasta)

    # Resumen agrupado del forecast por servicio
    fore_summary = df_fore_rango.groupby(["servicio", "tipo_mundo"]).agg({
        "traffic_forecast": "sum",
        "minutos_req": "sum",
        "meta_aht_plana": "first",
        "meta_ns": "first"
    }).reset_index()

    # FTEs requeridos en base a jornada estándar de 8 horas (480 minutos * num_dias)
    fore_summary["fte_dia_requerido"] = (fore_summary["minutos_req"] / (480.0 * num_dias)).round(1)

    # Unir Forecast y Real (Presencia)
    if not df_pres_rango.empty:
        matriz = pd.merge(fore_summary, df_pres_rango, on="servicio", how="left").fillna(0.0)
    else:
        matriz = fore_summary.copy()
        matriz["min_conectado"] = 0.0
        matriz["min_pausas"] = 0.0
        matriz["min_disponible"] = 0.0
        matriz["pct_auxiliares_real"] = 0.0
        matriz["fte_reales_conectados"] = 0.0
        matriz["fte_reales_disponibles"] = 0.0

    # Unir Métricas de GTR (Tráfico, AHT, NS)
    if not df_gtr_rango.empty:
        matriz = pd.merge(matriz, df_gtr_rango, on="servicio", how="left")
    else:
        matriz["trafico_real"] = np.nan
        matriz["aht_real_seg"] = np.nan
        matriz["ns_real"] = np.nan
        matriz["tHandle_sum"] = 0.0
        matriz["tHandle_count"] = 0
        matriz["sl_numerator"] = 0
        matriz["sl_denominator"] = 0

    # Aplicar filtro por mundo
    if filtro_mundo == "📞 Línea / Inbound Voz":
        matriz = matriz[matriz["tipo_mundo"].str.contains("Línea|Voz|Multi", case=False, na=False)]
    elif filtro_mundo == "💬 Canales Digitales":
        matriz = matriz[matriz["tipo_mundo"].str.contains("Digital|Multi", case=False, na=False)]
    elif filtro_mundo == "📂 Back Office":
        matriz = matriz[matriz["tipo_mundo"].str.contains("Back|BO|Multi", case=False, na=False)]

    if matriz.empty:
        st.info("Sin registros para los filtros seleccionados.")
        return

    # Construir filas ejecutivas con cálculos
    filas_tabla = []
    for _, row in matriz.iterrows():
        srv = row["servicio"]
        tipo = row["tipo_mundo"]
        fte_req = row["fte_dia_requerido"]
        fte_con = row["fte_reales_conectados"]
        fte_disp = row["fte_reales_disponibles"]
        aux_real = row["pct_auxiliares_real"]
        meta_aht = row["meta_aht_plana"]
        meta_ns = row.get("meta_ns", 80.0)
        min_req = row["minutos_req"]
        min_disp = row["min_disponible"]
        min_con = row["min_conectado"]
        min_pau = row["min_pausas"]

        traf_plan = row.get("traffic_forecast", 0.0)
        traf_real = row.get("trafico_real")
        aht_real = row.get("aht_real_seg")
        ns_real = row.get("ns_real")

        pct_desv_trafico = round(((traf_real - traf_plan) / traf_plan * 100.0), 1) if pd.notna(traf_real) and traf_plan > 0 else np.nan
        pct_desv_aht = round(((aht_real - meta_aht) / meta_aht * 100.0), 1) if pd.notna(aht_real) and meta_aht > 0 else np.nan

        gap_fte = fte_con - fte_req
        pct_capacidad = (min_disp / min_req * 100.0) if min_req > 0 else 100.0

        # Pérdida de tiempo por exceder el 14% de auxiliares
        min_aux_permitidos = min_con * (META_AUXILIARES_OFICIAL / 100.0)
        min_exceso_aux = max(0.0, min_pau - min_aux_permitidos)
        horas_exceso_aux = min_exceso_aux / 60.0

        causa = diagnosticar_causa_raiz(
            gap_personas=gap_fte,
            aux_real=aux_real,
            aux_meta=META_AUXILIARES_OFICIAL,
            pct_capacidad=pct_capacidad
        )

        filas_tabla.append({
            "Servicio": srv,
            "Canal": tipo,
            "FTE Req": fte_req,
            "FTE Con": fte_con,
            "Brecha FTE": gap_fte,
            "% Aux Real": aux_real,
            "Meta Aux": META_AUXILIARES_OFICIAL,
            "Horas Fuga Aux": horas_exceso_aux,
            "Tráfico Plan": traf_plan,
            "Tráfico Real": traf_real if pd.notna(traf_real) else np.nan,
            "% Desv Tráfico": pct_desv_trafico,
            "Meta AHT (s)": int(meta_aht),
            "AHT Real (s)": aht_real if pd.notna(aht_real) else np.nan,
            "% Desv AHT": pct_desv_aht,
            "% NS": ns_real if pd.notna(ns_real) else np.nan,
            "Meta NS": meta_ns,
            "% Capacidad": pct_capacidad,
            "Min. Requeridos": min_req,
            "Min. Disponibles": min_disp,
            "Min. Pausas": min_pau,
            "tHandle_sum": row.get("tHandle_sum", 0.0),
            "tHandle_count": row.get("tHandle_count", 0),
            "sl_numerator": row.get("sl_numerator", 0),
            "sl_denominator": row.get("sl_denominator", 0),
            "Diagnóstico Operativo": causa
        })

    df_ejecutiva = pd.DataFrame(filas_tabla).sort_values(by=["FTE Req"], ascending=False)

    # 4. Scorecard Macro de la Operación
    tot_fte_req = df_ejecutiva["FTE Req"].sum()
    tot_fte_con = df_ejecutiva["FTE Con"].sum()
    tot_min_req = df_ejecutiva["Min. Requeridos"].sum()
    tot_min_disp = df_ejecutiva["Min. Disponibles"].sum()
    tot_min_pau = df_ejecutiva["Min. Pausas"].sum()
    tot_min_con = tot_min_disp + tot_min_pau

    cumpl_global = (tot_min_disp / tot_min_req * 100.0) if tot_min_req > 0 else 0.0
    pct_aux_global = (tot_min_pau / tot_min_con * 100.0) if tot_min_con > 0 else 0.0
    gap_fte_global = tot_fte_con - tot_fte_req
    tot_horas_fuga = df_ejecutiva["Horas Fuga Aux"].sum()
    fte_fuga_equivalentes = tot_horas_fuga / (8.0 * num_dias)

    # Métricas consolidadas macro de GTR (NS y AHT ponderados)
    tot_sl_num = df_ejecutiva["sl_numerator"].sum()
    tot_sl_den = df_ejecutiva["sl_denominator"].sum()
    ns_global = (tot_sl_num / tot_sl_den * 100.0) if tot_sl_den > 0 else np.nan

    tot_th_sum = df_ejecutiva["tHandle_sum"].sum()
    tot_th_cnt = df_ejecutiva["tHandle_count"].sum()
    aht_global_seg = ((tot_th_sum / tot_th_cnt) / 1000.0) if tot_th_cnt > 0 else np.nan

    st.markdown("---")
    m1, m2, m3, m4, m5, m6 = st.columns(6)
    with m1:
        color_cumpl = "normal" if cumpl_global >= 90 else "inverse"
        st.metric(
            "Capacidad Neta",
            f"{cumpl_global:.1f}%",
            delta=f"{cumpl_global - 100.0:+.1f}% vs Req",
            delta_color=color_cumpl,
            help="% de minutos productivos reales frente al total requerido del mes."
        )
    with m2:
        color_gap = "normal" if gap_fte_global >= 0 else "inverse"
        delta_gap = f"{gap_fte_global:+.1f} FTEs" if num_dias == 1 else f"{gap_fte_global:+.1f} FTEs/día"
        st.metric(
            "Balance FTEs",
            f"{tot_fte_con:.1f} / {tot_fte_req:.1f}",
            delta=delta_gap,
            delta_color=color_gap,
            help="Asesores equivalentes conectados vs asesores requeridos."
        )
    with m3:
        color_aux = "normal" if pct_aux_global <= META_AUXILIARES_OFICIAL else "inverse"
        st.metric(
            "% Auxiliares",
            f"{pct_aux_global:.1f}%",
            delta=f"{pct_aux_global - META_AUXILIARES_OFICIAL:+.1f}% vs Meta (14%)",
            delta_color=color_aux,
            help="% del tiempo conectado consumido en pausas en toda la operación."
        )
    with m4:
        delta_fuga = f"≈ {fte_fuga_equivalentes:.1f} Asesores" if num_dias == 1 else f"≈ {fte_fuga_equivalentes:.1f} FTEs/día"
        st.metric(
            "Fuga Auxiliares",
            f"{tot_horas_fuga:.1f} h",
            delta=delta_fuga,
            delta_color="off",
            help="Horas hombre netas destruidas por haber superado el 14% de auxiliares."
        )
    with m5:
        if pd.notna(ns_global):
            color_ns = "normal" if ns_global >= 80.0 else "inverse"
            delta_ns = f"{ns_global - 80.0:+.1f}% vs 80%"
            st.metric(
                "Nivel de Servicio",
                f"{ns_global:.1f}%",
                delta=delta_ns,
                delta_color=color_ns,
                help="Nivel de Servicio consolidado (80/20) de todas las llamadas/chats atendidos en Genesys."
            )
        else:
            st.metric("Nivel de Servicio", "N/A", help="No aplica o sin datos en canales backoffice")
    with m6:
        if pd.notna(aht_global_seg):
            st.metric(
                "AHT Real Promedio",
                f"{aht_global_seg:.0f} s",
                delta=f"{aht_global_seg/60.0:.1f} min",
                delta_color="off",
                help="Tiempo medio de operación (Handle Time) real ponderado en Genesys."
            )
        else:
            st.metric("AHT Real Promedio", "N/A", help="No aplica")

    st.markdown("---")

    # 5. Árbol de Atribución y Descomposición de Capacidad (Suma y Resta en Horas y %)
    st.markdown("#### 🌳 Árbol de Atribución y Descomposición de Capacidad")
    st.caption("Explica de forma transparente cuánto sumó la dotación/conexión y cuánto restaron las pausas para llegar a la capacidad neta final (en horas y porcentaje).")

    col_ctrl_1, col_ctrl_2 = st.columns([2.2, 1.8])
    with col_ctrl_1:
        opciones_alcance = ["🌐 Consolidado Global (Toda la Operación)"] + [f"🔍 {s}" for s in df_ejecutiva["Servicio"].tolist()]
        sel_alcance = st.selectbox(
            "Alcance del Análisis de Capacidad:",
            opciones_alcance,
            index=0,
            help="Elige si deseas ver el balance de toda la operación consolidada o hacer zoom en un servicio específico."
        )
    with col_ctrl_2:
        modo_eje = st.radio(
            "Eje Principal del Gráfico:",
            ["Horas Equivalentes (h)", "Porcentaje de Capacidad (%)"],
            horizontal=True,
            help="Ambas magnitudes se visualizan simultáneamente en las etiquetas de las barras."
        )

    # Determinar métricas según alcance
    if sel_alcance.startswith("🌐"):
        w_nombre = f"Consolidado Global ({filtro_mundo})"
        w_min_req = tot_min_req
        w_min_disp = tot_min_disp
        w_min_pau = tot_min_pau
        w_min_con = tot_min_con
        w_fte_req = tot_fte_req
        w_fte_con = tot_fte_con
        w_hfuga = tot_horas_fuga
        w_meta_aht = None
    else:
        srv_limpio = sel_alcance.replace("🔍 ", "")
        w_nombre = srv_limpio
        fila_s = df_ejecutiva[df_ejecutiva["Servicio"] == srv_limpio].iloc[0]
        w_min_req = fila_s["Min. Requeridos"]
        w_min_disp = fila_s["Min. Disponibles"]
        w_min_pau = fila_s["Min. Pausas"]
        w_min_con = w_min_disp + w_min_pau
        w_fte_req = fila_s["FTE Req"]
        w_fte_con = fila_s["FTE Con"]
        w_hfuga = fila_s["Horas Fuga Aux"]
        w_meta_aht = fila_s["Meta AHT (s)"]

    w_h_req = w_min_req / 60.0
    w_h_disp = w_min_disp / 60.0
    w_h_pau = w_min_pau / 60.0
    w_h_con = w_min_con / 60.0
    w_delta_con_h = w_h_con - w_h_req

    w_pct_delta_con = (w_delta_con_h / w_h_req * 100.0) if w_h_req > 0 else 0.0
    w_pct_pau = -(w_h_pau / w_h_req * 100.0) if w_h_req > 0 else 0.0
    w_pct_disp = (w_h_disp / w_h_req * 100.0) if w_h_req > 0 else 0.0

    # Desglose de pausas: dentro de meta oficial (14%) vs exceso de auxiliares
    w_min_pau_meta = w_min_con * (META_AUXILIARES_OFICIAL / 100.0)
    w_min_pau_fuga = max(0.0, w_min_pau - w_min_pau_meta)
    w_h_pau_meta = w_min_pau_meta / 60.0
    w_h_pau_fuga = w_min_pau_fuga / 60.0
    w_pct_pau_meta = -(w_h_pau_meta / w_h_req * 100.0) if w_h_req > 0 else 0.0
    w_pct_pau_fuga = -(w_h_pau_fuga / w_h_req * 100.0) if w_h_req > 0 else 0.0

    # Métricas de GTR para el alcance seleccionado en el Árbol
    if sel_alcance.startswith("🌐"):
        w_traf_plan = df_ejecutiva["Tráfico Plan"].sum()
        w_traf_real = df_ejecutiva["Tráfico Real"].sum(skipna=True)
        w_aht_real = aht_global_seg
        w_ns_real = ns_global
        w_meta_ns = 80.0
    else:
        w_traf_plan = fila_s.get("Tráfico Plan", 0.0)
        w_traf_real = fila_s.get("Tráfico Real", np.nan)
        w_aht_real = fila_s.get("AHT Real (s)", np.nan)
        w_ns_real = fila_s.get("% NS", np.nan)
        w_meta_ns = fila_s.get("Meta NS", 80.0)

    # Cálculo del impacto de Demanda y AHT en Horas y Porcentaje para el Waterfall
    # 1. Impacto AHT: si AHT real > meta plana, destruye horas. (Meta - Real) * volumen / 3600
    if pd.notna(w_aht_real) and w_meta_aht and w_meta_aht > 0 and pd.notna(w_traf_real) and w_traf_real > 0:
        w_h_delta_aht = ((w_meta_aht - w_aht_real) * w_traf_real) / 3600.0
    elif pd.notna(w_aht_real) and w_meta_aht and w_meta_aht > 0 and w_traf_plan > 0:
        w_h_delta_aht = ((w_meta_aht - w_aht_real) * w_traf_plan) / 3600.0
    else:
        w_h_delta_aht = 0.0

    w_pct_delta_aht = (w_h_delta_aht / w_h_req * 100.0) if w_h_req > 0 else 0.0

    # 2. Impacto Demanda / Tráfico: (Plan - Real) * Meta_AHT / 3600
    # Si entró sobre-demanda (Real > Plan), representa una sobre-exigencia (resta horas de holgura / capacidad)
    if pd.notna(w_traf_real) and w_traf_plan > 0:
        ref_aht_para_vol = w_meta_aht if (w_meta_aht and w_meta_aht > 0) else (w_aht_real if pd.notna(w_aht_real) else 800.0)
        w_h_delta_demanda = ((w_traf_plan - w_traf_real) * ref_aht_para_vol) / 3600.0
    else:
        w_h_delta_demanda = 0.0

    w_pct_delta_demanda = (w_h_delta_demanda / w_h_req * 100.0) if w_h_req > 0 else 0.0

    # Capacidad efectiva ajustada por Demanda y AHT
    w_h_cap_efectiva = max(0.0, w_h_disp + w_h_delta_aht + w_h_delta_demanda)
    w_pct_cap_efectiva = (w_h_cap_efectiva / w_h_req * 100.0) if w_h_req > 0 else 0.0

    col_wat, col_diag = st.columns([1.55, 1.05])
    with col_wat:
        if modo_eje == "Porcentaje de Capacidad (%)":
            y_vals = [100.0, w_pct_delta_con, w_pct_pau, w_pct_disp, w_pct_delta_demanda, w_pct_delta_aht, w_pct_cap_efectiva]
            eje_y_lbl = "% de Capacidad Requerida"
        else:
            y_vals = [w_h_req, w_delta_con_h, -w_h_pau, w_h_disp, w_h_delta_demanda, w_h_delta_aht, w_h_cap_efectiva]
            eje_y_lbl = "Horas-Hombre Equivalentes"

        # Etiquetas claras e inequívocas para cada barra del Waterfall
        # Para Demanda: mostrar tanto la sobrecarga en llamadas (+X% volumen) como su impacto en capacidad (-Y%)
        if pd.notna(w_traf_real) and w_traf_plan > 0:
            pct_vol_diff = ((w_traf_real - w_traf_plan) / w_traf_plan * 100.0)
            lbl_demanda = f"<b>{w_h_delta_demanda:+,.1f} h</b><br>{w_pct_delta_demanda:+.1f}% Cap<br><span style='font-size:10px; color:#475569;'>({pct_vol_diff:+.1f}% Vol)</span>"
        else:
            lbl_demanda = "<b>0.0 h</b><br>0.0%"

        # Para AHT: mostrar tanto el desvío en segundos (+X% AHT) como su impacto en capacidad (-Y%)
        if pd.notna(w_aht_real) and w_meta_aht and w_meta_aht > 0:
            pct_aht_diff = ((w_aht_real - w_meta_aht) / w_meta_aht * 100.0)
            lbl_aht = f"<b>{w_h_delta_aht:+,.1f} h</b><br>{w_pct_delta_aht:+.1f}% Cap<br><span style='font-size:10px; color:#475569;'>({pct_aht_diff:+.1f}% AHT)</span>"
        else:
            lbl_aht = "<b>0.0 h</b><br>0.0%"

        text_vals = [
            f"<b>{w_h_req:,.1f} h</b><br>100.0%",
            f"<b>{w_delta_con_h:+,.1f} h</b><br>{w_pct_delta_con:+.1f}%",
            f"<b>{-w_h_pau:+,.1f} h</b><br>{w_pct_pau:+.1f}%",
            f"<b>{w_h_disp:,.1f} h</b><br>{w_pct_disp:.1f}%",
            lbl_demanda,
            lbl_aht,
            f"<b>{w_h_cap_efectiva:,.1f} h</b><br>{w_pct_cap_efectiva:.1f}%"
        ]

        color_final = "#10b981" if w_pct_cap_efectiva >= 95.0 else ("#f59e0b" if w_pct_cap_efectiva >= 85.0 else "#ef4444")

        fig_wat = go.Figure(go.Waterfall(
            orientation="v",
            measure=["absolute", "relative", "relative", "subtotal", "relative", "relative", "total"],
            x=[
                "1. Requerido Plan",
                "2. Asistencia / FTEs",
                "3. Pausas / Aux",
                "4. Disponible Real",
                "5. Efecto Demanda",
                "6. Efecto AHT",
                "7. Capacidad Neta Efectiva"
            ],
            y=y_vals,
            text=text_vals,
            textposition="outside",
            connector={"line": {"color": "#cbd5e1", "width": 1.5}},
            decreasing={"marker": {"color": "#ef4444"}},
            increasing={"marker": {"color": "#10b981"}},
            totals={"marker": {"color": color_final}}
        ))

        fig_wat.update_layout(
            title=f"Árbol de Atribución Integral — {w_nombre}",
            waterfallgap=0.25,
            margin=dict(l=20, r=20, t=50, b=20),
            yaxis=dict(title=eje_y_lbl),
            plot_bgcolor="rgba(0,0,0,0)",
            paper_bgcolor="rgba(0,0,0,0)"
        )
        st.plotly_chart(fig_wat, use_container_width=True)
        st.caption("💡 **Regla de Signos Operativa:** En Demanda y AHT, un incremento en el indicador operativo (+ llamadas o + segundos) **resta capacidad efectiva (barra roja que desciende)** porque sobrecarga o frena la operación.")

    with col_diag:
        es_superavit_con = w_delta_con_h >= 0
        es_cumplido = w_pct_cap_efectiva >= 85.0
        bg_card = "#f0fdf4" if w_pct_cap_efectiva >= 95.0 else ("#fffbeb" if w_pct_cap_efectiva >= 85.0 else "#fef2f2")
        border_card = "#10b981" if w_pct_cap_efectiva >= 95.0 else ("#f59e0b" if w_pct_cap_efectiva >= 85.0 else "#ef4444")

        # Texto del impacto de AHT
        if pd.notna(w_aht_real) and w_meta_aht and w_meta_aht > 0:
            diff_aht_s = w_aht_real - w_meta_aht
            pct_diff_aht = (diff_aht_s / w_meta_aht * 100.0)
            if diff_aht_s > 0:
                texto_linea_aht = f"<span style='color: #b91c1c; font-weight: 600;'>{w_pct_delta_aht:+.1f}%</span> ({w_h_delta_aht:+.1f} h) • AHT: {w_aht_real:.0f}s vs {w_meta_aht:.0f}s (+{pct_diff_aht:.1f}%)"
            else:
                texto_linea_aht = f"<span style='color: #15803d; font-weight: 600;'>{w_pct_delta_aht:+.1f}%</span> ({w_h_delta_aht:+.1f} h) • AHT: {w_aht_real:.0f}s vs {w_meta_aht:.0f}s ({pct_diff_aht:.1f}%)"
        else:
            texto_linea_aht = "Sin impacto (No aplica / Back Office)"

        # Texto del impacto de Tráfico
        if pd.notna(w_traf_real) and w_traf_plan > 0:
            diff_traf = w_traf_real - w_traf_plan
            pct_diff_traf = (diff_traf / w_traf_plan * 100.0)
            if diff_traf > 0:
                texto_linea_traf = f"<span style='color: #b91c1c; font-weight: 600;'>{w_pct_delta_demanda:+.1f}%</span> ({w_h_delta_demanda:+.1f} h) • Sobre-demanda (+{pct_diff_traf:.1f}%)"
            else:
                texto_linea_traf = f"<span style='color: #15803d; font-weight: 600;'>{w_pct_delta_demanda:+.1f}%</span> ({w_h_delta_demanda:+.1f} h) • Menor volumen ({pct_diff_traf:.1f}%)"
        else:
            texto_linea_traf = "Sin impacto medible de colas"

        # Texto del impacto de Nivel de Servicio
        if pd.notna(w_ns_real):
            color_ns_txt = "#15803d" if w_ns_real >= w_meta_ns else ("#b45309" if w_ns_real >= (w_meta_ns - 10.0) else "#b91c1c")
            texto_linea_ns = f"<span style='font-size: 15px; font-weight: 700; color: {color_ns_txt};'>{w_ns_real:.1f}%</span> (Meta: {w_meta_ns:.0f}%)"
        else:
            texto_linea_ns = "N/A"

        st.markdown(
            f"""
            <div style="background: {bg_card}; border: 1px solid {border_card}; border-radius: 10px; padding: 14px 16px; margin-top: 10px;">
                <div style="font-weight: 700; font-size: 14.5px; color: #0f172a; margin-bottom: 8px;">
                    ⚖️ Descomposición de Fuerzas: ¿Qué sumó y qué restó?
                </div>
                <div style="font-size: 12.5px; color: #334155; line-height: 1.6;">
                    • <b>1. Requerido Plan (100%):</b> <code>100.0%</code> ({w_h_req:,.1f} h | {w_fte_req:.1f} FTEs{'/día' if num_dias > 1 else ''})<br>
                    • <b>2. Conexión / Asistencia:</b> <span style="color: {'#15803d' if es_superavit_con else '#b91c1c'}; font-weight: 600;">{w_pct_delta_con:+.1f}%</span> ({w_delta_con_h:+,.1f} h)<br>
                    • <b>3. Pausas / Auxiliares:</b> <span style="color: #b91c1c; font-weight: 600;">{w_pct_pau:.1f}%</span> ({-w_h_pau:.1f} h | Fuga: {w_pct_pau_fuga:.1f}%)<br>
                    • <b>4. Disponible en Genesys:</b> <b>{w_pct_disp:.1f}%</b> ({w_h_disp:,.1f} h)<br>
                    <hr style="margin: 6px 0; border: none; border-top: 1px dashed #cbd5e1;">
                    • <b>5. Presión de Demanda:</b> {texto_linea_traf}<br>
                    • <b>6. Desvío de AHT (Eficiencia):</b> {texto_linea_aht}<br>
                    <hr style="margin: 6px 0; border: none; border-top: 1px dashed #cbd5e1;">
                    • <b>7. Capacidad Neta Efectiva:</b> <span style="font-size: 14px; font-weight: 700; color: {'#15803d' if es_cumplido else '#b91c1c'};">{w_pct_cap_efectiva:.1f}%</span> ({w_h_cap_efectiva:,.1f} h)<br>
                    • <b>🎯 Nivel de Servicio Resultante:</b> {texto_linea_ns}
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        # Diagnóstico narrativo de causa raíz
        if w_pct_cap_efectiva >= 95.0 and (pd.isna(w_ns_real) or w_ns_real >= w_meta_ns):
            st.success(f"✅ **Operación Cumplida y Blindada:** La capacidad efectiva cubrió el **{w_pct_cap_efectiva:.1f}%** del requerimiento y el Nivel de Servicio cerró en meta ({w_ns_real:.1f}%).")
        elif pd.notna(w_ns_real) and w_ns_real < w_meta_ns and w_pct_delta_aht < -5.0:
            st.error(
                f"🚨 **Déficit por AHT Desbordado:** El aumento en los tiempos de llamada restó **{abs(w_h_delta_aht):.1f} horas de capacidad** ({w_pct_delta_aht:.1f}%), siendo el factor determinante que hundió el NS al **{w_ns_real:.1f}%**."
            )
        elif pd.notna(w_ns_real) and w_ns_real < w_meta_ns and w_pct_delta_demanda < -5.0:
            st.error(
                f"🚨 **Colapso por Sobre-Demanda:** La llegada de volumen no previsto restó **{abs(w_h_delta_demanda):.1f} horas de holgura** ({w_pct_delta_demanda:.1f}%), sobrepasando la capacidad planificada."
            )
        elif not es_superavit_con:
            st.error(
                f"🚨 **Déficit por Falta de Conexión:** Faltaron **{abs(w_delta_con_h):.1f} horas** de personal ({w_pct_delta_con:.1f}% vs plan) para sostener la operación."
            )
        elif es_superavit_con and w_pct_pau_fuga < -5.0:
            st.warning(
                f"🟠 **Fuga en Auxiliares:** Se contó con suficiente personal ({w_pct_delta_con:+.1f}%), pero las pausas no autorizadas destruyeron **{w_h_pau_fuga:.1f} horas**."
            )
        else:
            st.warning(f"⚠️ **Capacidad Ajustada:** Cumplimiento efectivo del **{w_pct_cap_efectiva:.1f}%** frente a la exigencia planificada.")

    st.markdown("---")

    # 6. Tabla Matriz Ejecutiva Panorámica
    st.markdown("#### 📊 Matriz Panorámica de Cumplimiento y Causa Raíz")
    st.caption("Visión gerencial integral: Requerido vs Conectados vs Auxiliares vs Tráfico vs AHT vs Nivel de Servicio (NS 80/20). Haz clic en cualquier servicio para ver su detalle.")

    def estilo_gap(val):
        if pd.isna(val): return ""
        if val >= 0: color = "#10b981"
        elif val >= -2.0: color = "#f59e0b"
        else: color = "#ef4444"
        return f"background-color: {color}20; color: {color}; font-weight: 700;"

    def estilo_cumpl(val):
        if pd.isna(val): return ""
        if val >= 95.0: color = "#10b981"
        elif val >= 85.0: color = "#f59e0b"
        else: color = "#ef4444"
        return f"background-color: {color}20; color: {color}; font-weight: 700;"

    def estilo_aux(val):
        if pd.isna(val) or val == 0: return ""
        if val <= META_AUXILIARES_OFICIAL: color = "#10b981"
        elif val <= (META_AUXILIARES_OFICIAL + 4.0): color = "#f59e0b"
        else: color = "#ef4444"
        return f"background-color: {color}20; color: {color}; font-weight: 700;"

    def estilo_ns(val):
        if pd.isna(val): return ""
        if val >= 80.0: color = "#10b981"
        elif val >= 70.0: color = "#f59e0b"
        else: color = "#ef4444"
        return f"background-color: {color}20; color: {color}; font-weight: 700;"

    def estilo_desv_trafico(val):
        if pd.isna(val): return ""
        if val <= 5.0: color = "#10b981"
        elif val <= 15.0: color = "#f59e0b"
        else: color = "#ef4444"
        return f"background-color: {color}20; color: {color}; font-weight: 600;"

    cols_matriz_ejecutiva = [
        "Servicio", "Canal", "FTE Req", "FTE Con", "Brecha FTE",
        "% Aux Real", "Meta Aux", "Tráfico Plan", "Tráfico Real", "% Desv Tráfico",
        "Meta AHT (s)", "AHT Real (s)", "% NS", "% Capacidad", "Diagnóstico Operativo"
    ]

    styler_matriz = (
        df_ejecutiva[cols_matriz_ejecutiva].style
        .map(estilo_gap, subset=["Brecha FTE"])
        .map(estilo_cumpl, subset=["% Capacidad"])
        .map(estilo_aux, subset=["% Aux Real"])
        .map(estilo_ns, subset=["% NS"])
        .map(estilo_desv_trafico, subset=["% Desv Tráfico"])
    )

    evento = st.dataframe(
        styler_matriz,
        use_container_width=True,
        hide_index=True,
        on_select="rerun",
        selection_mode="single-row",
        key="tabla_matriz_capacidad_v2",
        column_config={
            "Servicio": st.column_config.TextColumn("Servicio", help="Nombre oficial en Genesys y Base del Requerido"),
            "Canal": st.column_config.TextColumn("Mundo / Canal"),
            "FTE Req": st.column_config.NumberColumn("FTE Req", format="%.1f", help="Asesores requeridos en el mes"),
            "FTE Con": st.column_config.NumberColumn("FTE Con", format="%.1f", help="Asesores conectados en Genesys"),
            "Brecha FTE": st.column_config.NumberColumn("Brecha FTE", format="%+.1f", help="Diferencia de personal: Conectados - Requeridos"),
            "% Aux Real": st.column_config.NumberColumn("% Aux", format="%.1f%%", help="% de tiempo en pausas"),
            "Meta Aux": st.column_config.NumberColumn("Meta Aux", format="%.0f%%"),
            "Tráfico Plan": st.column_config.NumberColumn("Tráfico Plan", format="%.0f", help="Volumen proyectado en forecast"),
            "Tráfico Real": st.column_config.NumberColumn("Tráfico Real", format="%.0f", help="Volumen real recibido en Genesys"),
            "% Desv Tráfico": st.column_config.NumberColumn("% Desv Tráfico", format="%+.1f%%", help="Desviación de volumen vs Forecast"),
            "Meta AHT (s)": st.column_config.NumberColumn("Meta AHT", format="%d s", help="Meta plana fija oficial"),
            "AHT Real (s)": st.column_config.NumberColumn("AHT Real", format="%.0f s", help="Handle time real promedio"),
            "% NS": st.column_config.NumberColumn("% NS", format="%.1f%%", help="Nivel de Servicio 80/20"),
            "% Capacidad": st.column_config.NumberColumn("% Capacidad", format="%.1f%%", help="Minutos disponibles ÷ Minutos requeridos"),
            "Diagnóstico Operativo": st.column_config.TextColumn("Causa Raíz / Veredicto"),
        }
    )

    # 7. Servicio Seleccionado para el Detalle Intradía
    filas_sel = evento.selection.get("rows", [])
    if filas_sel:
        idx_sel = filas_sel[0]
        srv_detalle = df_ejecutiva.iloc[idx_sel]["Servicio"]
        fila_detalle = df_ejecutiva.iloc[idx_sel]
    else:
        srv_detalle = df_ejecutiva.iloc[0]["Servicio"]
        fila_detalle = df_ejecutiva.iloc[0]

    # 8. Diagnóstico Narrativo Ejecutivo del Servicio Seleccionado
    st.markdown("---")
    st.subheader(f"🔎 Diagnóstico Detallado — {srv_detalle}")

    d_req = fila_detalle["FTE Req"]
    d_con = fila_detalle["FTE Con"]
    d_gap = fila_detalle["Brecha FTE"]
    d_aux = fila_detalle["% Aux Real"]
    d_cap = fila_detalle["% Capacidad"]
    d_aht = fila_detalle["Meta AHT (s)"]
    d_mreq = fila_detalle["Min. Requeridos"]
    d_mdisp = fila_detalle["Min. Disponibles"]
    d_mpau = fila_detalle["Min. Pausas"]
    d_hfuga = fila_detalle["Horas Fuga Aux"]

    texto_personas = (
        f"🟢 Se contó con dotación suficiente (**{d_con:.1f} FTEs** vs **{d_req:.1f} FTEs** requeridos, **{d_gap:+.1f}** de holgura)"
        if d_gap >= 0
        else f"🔴 Se presentó un déficit de personal de **{d_gap:+.1f} FTEs** (**{d_con:.1f}** conectados frente a **{d_req:.1f}** solicitados)"
    )

    texto_aux = (
        f"🟢 La disciplina de auxiliares estuvo controlada en un **{d_aux:.1f}%** (dentro de la meta oficial del 14%)"
        if d_aux <= META_AUXILIARES_OFICIAL
        else f"🟠 Se registró sobreconsumo de auxiliares con un **{d_aux:.1f}%** frente al 14% meta, destruyendo **{d_hfuga:.1f} horas de capacidad**"
    )

    d_traf_plan = fila_detalle.get("Tráfico Plan", 0.0)
    d_traf_real = fila_detalle.get("Tráfico Real", np.nan)
    d_desv_traf = fila_detalle.get("% Desv Tráfico", np.nan)
    d_aht_real = fila_detalle.get("AHT Real (s)", np.nan)
    d_desv_aht = fila_detalle.get("% Desv AHT", np.nan)
    d_ns = fila_detalle.get("% NS", np.nan)
    d_meta_ns = fila_detalle.get("Meta NS", 80.0)

    # Texto de demanda
    if pd.notna(d_traf_real) and d_traf_plan > 0:
        texto_demanda = (
            f"🔴 Se atendió una sobre-demanda de **{d_traf_real:,.0f} interacciones** vs **{d_traf_plan:,.0f} planificadas** (**{d_desv_traf:+.1f}%** de sobrecarga no prevista)"
            if d_desv_traf > 5.0
            else f"🟢 El volumen de interacciones estuvo alineado con el forecast (**{d_traf_real:,.0f}** recibidas vs **{d_traf_plan:,.0f}** proyectadas, **{d_desv_traf:+.1f}%**)"
        )
    else:
        texto_demanda = "ℹ️ Servicio sin medición de volumen de llamadas/chats en colas directas."

    # Texto de AHT
    if pd.notna(d_aht_real) and d_aht and d_aht > 0:
        texto_aht_diag = (
            f"🔴 El AHT real se desbordó a **{d_aht_real:.0f}s** frente a la meta plana de **{d_aht}s** (**{d_desv_aht:+.1f}%**), destruyendo capacidad operativa por interacción"
            if d_aht_real > d_aht
            else f"🟢 El AHT real estuvo eficiente en **{d_aht_real:.0f}s** (por debajo de la meta plana de **{d_aht}s**, **{d_desv_aht:+.1f}%**)"
        )
    elif pd.notna(d_aht_real):
        texto_aht_diag = f"ℹ️ AHT promedio registrado: **{d_aht_real:.0f} segundos**."
    else:
        texto_aht_diag = f"ℹ️ Evaluado con meta plana de **{d_aht} segundos**."

    # Texto de NS
    if pd.notna(d_ns):
        texto_ns_diag = (
            f"🟢 **{d_ns:.1f}%** (Cumpliendo la meta de servicio del {d_meta_ns:.0f}%)"
            if d_ns >= d_meta_ns
            else f"🔴 **{d_ns:.1f}%** (Por debajo del objetivo del {d_meta_ns:.0f}%)"
        )
    else:
        texto_ns_diag = "N/A (Back Office / Casos)"

    if num_dias == 1:
        texto_intro = f"Para el servicio <b>{srv_detalle}</b> el <b>{fecha_desde}</b>, la base del requerido del mes dimensionó una exigencia de <b>{d_mreq:,.0f} minutos hombre</b> ({d_req:.1f} FTEs)."
    else:
        texto_intro = f"Para el servicio <b>{srv_detalle}</b> en el periodo <b>del {fecha_desde} al {fecha_hasta}</b> ({num_dias} días evaluados), la base del requerido dimensionó una exigencia promedio de <b>{d_req:.1f} FTEs/día</b> ({d_mreq:,.0f} minutos hombre acumulados)."

    if "DREAM TEAM" in srv_detalle.upper() or "DT FFP" in srv_detalle.upper():
        st.info("ℹ️ **Operación Homologada Dream Team:** Se consolida la presencia del equipo multi-skill de Genesys con la exigencia unificada de Voz, Chat, WhatsApp y Casos Backoffice (con Casos Backoffice computado como estado productivo).")
    elif "AGENCIAS" in srv_detalle.upper():
        st.info("ℹ️ **Operación Homologada Agencias:** Se unifican los niveles operativos (N1 y N3) registrados en nómina contra la proyección de dimensionamiento.")

    st.markdown(
        f"""
        <div style="background: #ffffff; border-left: 5px solid {'#10b981' if d_cap >= 90 and (pd.isna(d_ns) or d_ns >= d_meta_ns) else '#ef4444'}; border-radius: 8px; padding: 14px 18px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); margin-bottom: 16px;">
            <span style="font-size: 15px; font-weight: 700; color: #0f172a;">Diagnóstico Integral: Capacidad, Eficiencia y Nivel de Servicio:</span><br>
            <p style="font-size: 13.5px; color: #334155; margin-top: 6px; line-height: 1.6;">
                {texto_intro}<br>
                • <b>1. Conexión / Asistencia:</b> {texto_personas}.<br>
                • <b>2. Auxiliares y Pausas:</b> {texto_aux}. De los minutos conectados, <b>{d_mpau:,.0f} minutos</b> se consumieron en pausas.<br>
                • <b>3. Capacidad Neta Disponible:</b> Quedaron <b>{d_mdisp:,.0f} minutos productivos</b>, alcanzando un <b>{d_cap:.1f}% de capacidad</b> frente al plan.<br>
                • <b>4. Tráfico y Demanda:</b> {texto_demanda}.<br>
                • <b>5. Desempeño AHT:</b> {texto_aht_diag}.<br>
                • <b>🎯 Resultado Nivel de Servicio (NS 80/20):</b> {texto_ns_diag}.
            </p>
        </div>
        """,
        unsafe_allow_html=True
    )

    # 9. Vista Detallada: Intradía (si 1 día) o Evolución Diaria + Zoom (si Rango)
    if num_dias > 1:
        tab_evol, tab_intra = st.tabs([
            f"📅 Evolución Diaria ({num_dias} días)",
            "🕒 Zoom Intradía por Intervalo (30 min)"
        ])
        with tab_evol:
            st.markdown(f"##### 📈 Curva de Capacidad Día a Día — {srv_detalle}")
            st.caption("Comportamiento diario de la capacidad requerida vs ejecutada durante las fechas evaluadas.")
            df_evol = calcular_evolucion_diaria_servicio(fecha_desde, fecha_hasta, srv_detalle, df_fore_all)
            if not df_evol.empty:
                fig_ev = go.Figure()
                fig_ev.add_trace(go.Scatter(
                    x=df_evol["fecha"],
                    y=df_evol["fte_req"],
                    mode="lines+markers",
                    name="1. Requerido del Mes (FTEs)",
                    line=dict(color="#f59e0b", width=3, dash="dash")
                ))
                fig_ev.add_trace(go.Scatter(
                    x=df_evol["fecha"],
                    y=df_evol["fte_con"],
                    mode="lines+markers",
                    name="2. Conectados Totales (FTEs)",
                    line=dict(color="#94a3b8", width=2)
                ))
                fig_ev.add_trace(go.Bar(
                    x=df_evol["fecha"],
                    y=df_evol["fte_disp"],
                    name="3. Disponible Efectivo (FTEs)",
                    marker_color="#2563eb",
                    opacity=0.75
                ))
                fig_ev.update_layout(
                    title=f"Evolución Diaria de Personal y Capacidad — {srv_detalle}",
                    xaxis=dict(title="Fecha", tickangle=-30),
                    yaxis=dict(title="Equivalente de Asesores (FTEs)"),
                    hovermode="x unified",
                    margin=dict(l=20, r=20, t=40, b=20),
                    legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
                )
                st.plotly_chart(fig_ev, use_container_width=True)

                with st.expander(f"📋 Ver Desglose Numérico Día a Día ({srv_detalle})", expanded=False):
                    st.dataframe(
                        df_evol.rename(columns={
                            "fecha": "Fecha",
                            "fte_req": "FTE Requerido",
                            "fte_con": "FTE Conectado",
                            "fte_disp": "FTE Disponible",
                            "brecha_fte": "Brecha FTE",
                            "pct_aux": "% Auxiliares",
                            "pct_capacidad": "% Capacidad",
                            "minutos_req": "Min. Requeridos",
                            "min_disponible": "Min. Disponibles"
                        })[[
                            "Fecha", "FTE Requerido", "FTE Conectado", "FTE Disponible",
                            "Brecha FTE", "% Auxiliares", "% Capacidad", "Min. Requeridos", "Min. Disponibles"
                        ]],
                        use_container_width=True,
                        hide_index=True,
                        column_config={
                            "Fecha": st.column_config.TextColumn("Fecha", width="small"),
                            "FTE Requerido": st.column_config.NumberColumn(format="%.1f"),
                            "FTE Conectado": st.column_config.NumberColumn(format="%.1f"),
                            "FTE Disponible": st.column_config.NumberColumn(format="%.1f"),
                            "Brecha FTE": st.column_config.NumberColumn(format="%+.1f"),
                            "% Auxiliares": st.column_config.NumberColumn(format="%.1f%%"),
                            "% Capacidad": st.column_config.NumberColumn(format="%.1f%%"),
                            "Min. Requeridos": st.column_config.NumberColumn(format="%.0f m"),
                            "Min. Disponibles": st.column_config.NumberColumn(format="%.0f m"),
                        }
                    )
            else:
                st.info("Sin registros de evolución diaria para este servicio en el rango seleccionado.")

        with tab_intra:
            dias_disp_zoom = sorted(df_fore_rango["fecha"].unique().tolist())
            col_dz, _ = st.columns([2.5, 3.5])
            with col_dz:
                dia_zoom = st.selectbox(
                    "📅 Selecciona un día para examinar sus 48 franjas horarias:",
                    dias_disp_zoom,
                    index=len(dias_disp_zoom) - 1,
                    key="cap_dia_zoom_picker"
                )
            _render_curva_y_tabla_intradia(dia_zoom, srv_detalle, df_fore_all)
    else:
        st.markdown("##### 📈 Curva Intradía de Cobertura (30 min × FTEs = Minutos)")
        st.caption("Compara en cada intervalo cuántas personas exigía la base del requerido del mes vs cuántas estaban conectadas y cuántas efectivamente disponibles.")
        _render_curva_y_tabla_intradia(fecha_desde, srv_detalle, df_fore_all)

